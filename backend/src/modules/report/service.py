from datetime import datetime
from io import BytesIO

import openpyxl
import openpyxl.styles
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from src.common.enums.document_status import DocumentStatus
from src.common.schemas.token_user_info import TokenUserInfo
from src.modules.documents.repositories import DocumentRepository


class ReportService:
    def __init__(self, db: AsyncSession, doc_repo: DocumentRepository):
        self.db = db
        self.doc_repo = doc_repo

    async def get(
        self,
        user: TokenUserInfo,
        start: datetime,
        end: datetime,
    ):
        STATUSES_MAP = {
            DocumentStatus.APPROVED: "Зачёт",
            DocumentStatus.REJECTED: "Незачёт",
        }

        headers = ["ФИО студента", "Группа", "Тема работы", "Статус", "Оценка"]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Отчёт"

        header_font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        header_fill = openpyxl.styles.PatternFill("solid", fgColor="155dfc")
        header_alignment = openpyxl.styles.Alignment(
            horizontal="center", vertical="center"
        )

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        docs = await self.doc_repo.get_all_processed(
            user_id=user.id, from_date=start, to_date=end
        )
        row_idx = 2
        for doc in docs:
            for student in doc.authors:
                ws.cell(
                    row=row_idx,
                    column=1,
                    value=f"{student.last_name} {student.first_name} {student.middle_name or ''}",
                )
                ws.cell(
                    row=row_idx,
                    column=2,
                    value=(student.group or "Группа не распознана"),
                )
                ws.cell(row=row_idx, column=3, value=doc.topic)
                ws.cell(
                    row=row_idx,
                    column=4,
                    value=STATUSES_MAP.get(doc.status, "Ошибка"),
                )
                ws.cell(row=row_idx, column=5, value=doc.score)

                # выравнивание
                for col_idx in range(1, 6):
                    ws.cell(
                        row=row_idx, column=col_idx
                    ).alignment = openpyxl.styles.Alignment(
                        horizontal="left", vertical="top", wrap_text=True
                    )

                row_idx += 1

        column_widths = [30, 15, 40, 15, 10]
        for i, width in enumerate(column_widths, start=1):
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(i)
            ].width = width

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": 'attachment; filename="report.xlsx"'
            },
        )
